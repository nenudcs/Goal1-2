"""Strict image manifests and leakage-aware splits for the detection tasks.

Run this module and its tests only in the competition container. Image paths are
indexed once, independently of labels; only image arrays reach model adapters.
"""
import csv
import itertools
import json
import math
import random
import re
from collections import Counter, defaultdict
from pathlib import Path

from .common import file_hash, fingerprint


LABEL_MAP = {"true": 0, "fake": 1, "compositing": 2, "composition": 2}
SPLIT_NAMES = ("train", "val", "calibration")
_COLUMNS = {
    "accession": ("accessionnumber", "accession", "caseid"),
    "series_uid": ("seriesuid", "seriesinstanceuid"),
    "patient_id": ("patientid",),
    "a": ("srcimg", "a", "accessiona", "srcaccessionnumber"),
    "b": ("descimg", "b", "accessionb", "dstaccessionnumber"),
    "label": ("label", "y", "isduplicate"),
    "split": ("split", "subset"),
}


def _text_id(value, context):
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"{context}: IDs must be nonempty text, including in XLSX; "
                         "numeric cells can irreversibly lose precision")
    value = value.strip()
    from .submission import identifier
    identifier(value)
    if re.fullmatch(r"[+-]?\d+(?:\.\d+)?[eE][+-]?\d+", value):
        raise ValueError(f"{context}: scientific-notation ID is ambiguous: {value!r}")
    return value


def _table(path):
    """Read typed XLSX cells or literal CSV text without pandas ID inference."""
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(path)
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            rows = []
            for row in workbook.active.iter_rows():
                if any(cell.data_type == "f" for cell in row):
                    raise ValueError(f"{path}: formulas are not accepted in label/ID tables")
                rows.append([cell.value for cell in row])
        finally:
            workbook.close()
    elif path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.reader(stream))
    else:
        raise ValueError(f"{path}: expected .xlsx or .csv")
    rows = [row for row in rows if any(value not in (None, "") for value in row)]
    if not rows:
        raise ValueError(f"{path}: empty table")
    while rows[0] and rows[0][-1] in (None, ""):
        rows[0].pop()
    headers = [re.sub(r"[\W_]+", "", str(value).casefold()) if value is not None else "" for value in rows[0]]
    if any(not value for value in headers) or len(set(headers)) != len(headers):
        raise ValueError(f"{path}: empty or duplicate column names")
    output = []
    for number, row in enumerate(rows[1:], 2):
        if any(value not in (None, "") for value in row[len(headers):]):
            raise ValueError(f"{path}:{number}: too many cells")
        row = row[:len(headers)]
        output.append(dict(zip(headers, row + [None] * (len(headers) - len(row)))))
    return output


def _cell(row, field, context, required=True):
    columns = [key for key in _COLUMNS[field] if key in row]
    if len(columns) > 1:
        raise ValueError(f"{context}: ambiguous columns for {field}: {columns}")
    value = row[columns[0]] if columns else None
    if required and value in (None, ""):
        raise ValueError(f"{context}: missing {field}")
    return value


def _label_file(directory, basename):
    candidates = [directory / f"{basename}{suffix}" for suffix in (".xlsx", ".csv")]
    existing = [path for path in candidates if path.is_file()]
    if len(existing) != 1:
        raise ValueError(f"Expected exactly one XLSX/CSV label table for {basename}: {existing}")
    return existing[0]


def _canonical_array(image, original_affine, original_shape, extra):
    import nibabel as nib
    import numpy as np
    affine = np.asarray(image.affine, dtype=np.float64)
    if (affine.shape != (4, 4) or not np.isfinite(affine).all()
            or abs(np.linalg.det(affine[:3, :3])) < 1e-12
            or not np.allclose(affine[3], [0, 0, 0, 1])):
        raise ValueError("Image has invalid or singular physical geometry")
    if len(image.shape) != 3 or any(size < 1 for size in image.shape):
        raise ValueError(f"Expected a nonempty 3D scalar image, got {image.shape}")
    # Canonicalization permutes/flips voxels; it never resamples to a fixed cube.
    canonical = nib.as_closest_canonical(image)
    xyz = canonical.get_fdata(dtype=np.float32)
    if not np.isfinite(xyz).all():
        raise ValueError("Image contains NaN or infinite voxels")
    dhw = np.ascontiguousarray(xyz.transpose(2, 1, 0), dtype=np.float32)
    metadata = {
        "shape_dhw": list(dhw.shape), "original_shape_xyz": list(original_shape),
        "original_affine_ras_xyz": np.asarray(original_affine).tolist(),
        "affine_ras_xyz": canonical.affine.tolist(),
        "spacing_dhw": list(map(float, nib.affines.voxel_sizes(canonical.affine)[::-1])),
        "orientation": list(nib.aff2axcodes(canonical.affine)),
        "array_order": "DHW; affine indexes XYZ = WHD", **extra,
    }
    return dhw, metadata


def _dicom_header(path):
    import pydicom
    header = pydicom.dcmread(str(path), stop_before_pixels=True)
    # DICOM segmentations and radiotherapy objects are annotations, not inputs.
    if str(getattr(header, "Modality", "")) in {"SEG", "RTSTRUCT", "RTDOSE", "RTPLAN"}:
        return None
    if "Segmentation" in getattr(getattr(header, "SOPClassUID", None), "name", ""):
        return None
    for field in ("AccessionNumber", "SeriesInstanceUID", "SOPInstanceUID"):
        _text_id(getattr(header, field, None), f"{path}: {field}")
    if int(getattr(header, "NumberOfFrames", 1)) != 1:
        raise ValueError(f"{path}: multi-frame DICOM is unsupported; convert explicitly first")
    if (int(getattr(header, "SamplesPerPixel", 1)) != 1
            or str(getattr(header, "PhotometricInterpretation", "")) not in {"MONOCHROME1", "MONOCHROME2"}):
        raise ValueError(f"{path}: only scalar monochrome DICOM is supported")
    if int(getattr(header, "Rows", 0)) < 1 or int(getattr(header, "Columns", 0)) < 1:
        raise ValueError(f"{path}: DICOM has no valid image dimensions")
    return header


def _load_dicom(record):
    import nibabel as nib
    import numpy as np
    import pydicom
    files = record.get("files", [])
    if not files or len(files) != len(set(files)):
        raise ValueError("DICOM record must contain distinct image files")
    headers = [_dicom_header(path) for path in files]
    if any(header is None for header in headers):
        raise ValueError("DICOM annotation object cannot be loaded as an image")
    first = headers[0]
    orientation = np.asarray(getattr(first, "ImageOrientationPatient", []), dtype=float)
    spacing = np.asarray(getattr(first, "PixelSpacing", []), dtype=float)
    if (orientation.shape != (6,) or not np.isfinite(orientation).all()
            or not np.allclose([np.linalg.norm(orientation[:3]), np.linalg.norm(orientation[3:])], 1, atol=1e-3)
            or abs(np.dot(orientation[:3], orientation[3:])) > 1e-3):
        raise ValueError("DICOM requires valid orthonormal ImageOrientationPatient")
    if spacing.shape != (2,) or not np.isfinite(spacing).all() or (spacing <= 0).any():
        raise ValueError("DICOM requires positive physical PixelSpacing")
    normal = np.cross(orientation[:3], orientation[3:])
    normal /= np.linalg.norm(normal)
    positions, sop_ids = [], set()
    patient = None
    for path, header in zip(files, headers):
        if (str(header.AccessionNumber).strip() != record["accession"]
                or str(header.SeriesInstanceUID).strip() != record["series_uid"]):
            raise ValueError(f"{path}: mixed accession or series in DICOM stack")
        identity = str(header.SOPInstanceUID)
        if identity in sop_ids:
            raise ValueError(f"{path}: duplicate SOPInstanceUID")
        sop_ids.add(identity)
        current_patient = str(getattr(header, "PatientID", "")).strip() or None
        if patient and current_patient and patient != current_patient:
            raise ValueError(f"{path}: mixed PatientID in DICOM stack")
        patient = patient or current_patient
        current_orientation = np.asarray(getattr(header, "ImageOrientationPatient", []), dtype=float)
        current_spacing = np.asarray(getattr(header, "PixelSpacing", []), dtype=float)
        position = np.asarray(getattr(header, "ImagePositionPatient", []), dtype=float)
        if (current_orientation.shape != (6,) or current_spacing.shape != (2,)
                or not np.allclose(current_orientation, orientation, atol=1e-4)
                or not np.allclose(current_spacing, spacing, atol=1e-4)
                or (header.Rows, header.Columns) != (first.Rows, first.Columns)
                or str(getattr(header, "Modality", "")) != str(getattr(first, "Modality", ""))
                or str(header.PhotometricInterpretation) != str(first.PhotometricInterpretation)):
            raise ValueError(f"{path}: mixed geometry or image type in DICOM series")
        if position.shape != (3,) or not np.isfinite(position).all():
            raise ValueError(f"{path}: valid ImagePositionPatient is required")
        positions.append(position)
    order = np.argsort(np.asarray(positions) @ normal)
    ordered_positions = np.asarray(positions)[order]
    if len(order) > 1:
        steps = np.diff(ordered_positions, axis=0)
        projections = steps @ normal
        if (projections <= 1e-4).any():
            raise ValueError("Repeated physical slice positions: mixed time/echo stacks are unsupported")
        step = np.median(steps, axis=0)
        if not np.allclose(steps, step, rtol=1e-3, atol=1e-2):
            raise ValueError("Irregular DICOM slice spacing is unsupported; convert/resample explicitly")
    else:
        thickness = float(getattr(first, "SpacingBetweenSlices", getattr(first, "SliceThickness", 0)))
        if not math.isfinite(thickness) or thickness <= 0:
            raise ValueError("Single-slice DICOM needs positive slice spacing/thickness")
        step = normal * thickness
    frames = []
    for index in order:
        dataset = pydicom.dcmread(str(files[int(index)]))
        pixels = dataset.pixel_array
        if pixels.shape != (int(first.Rows), int(first.Columns)):
            raise ValueError(f"{files[int(index)]}: invalid scalar pixel shape {pixels.shape}")
        # Do not silently approximate arbitrary modality LUTs with linear rescale.
        if getattr(dataset, "ModalityLUTSequence", None):
            raise ValueError("DICOM ModalityLUTSequence is unsupported; convert explicitly")
        slope = float(getattr(dataset, "RescaleSlope", 1))
        intercept = float(getattr(dataset, "RescaleIntercept", 0))
        if not math.isfinite(slope) or slope == 0 or not math.isfinite(intercept):
            raise ValueError("DICOM has invalid intensity rescaling")
        frames.append(pixels.astype(np.float32) * slope + intercept)
    xyz = np.stack(frames).transpose(2, 1, 0)
    lps_affine = np.eye(4)
    lps_affine[:3, 0] = orientation[:3] * spacing[1]
    lps_affine[:3, 1] = orientation[3:] * spacing[0]
    lps_affine[:3, 2] = step
    lps_affine[:3, 3] = ordered_positions[0]
    ras_affine = np.diag([-1., -1., 1., 1.]) @ lps_affine
    image = nib.Nifti1Image(xyz, ras_affine)
    return _canonical_array(image, ras_affine, xyz.shape, {
        "kind": "dicom", "patient_id": patient,
        "source_files_in_physical_order": [files[int(index)] for index in order],
        "photometric_interpretation": str(first.PhotometricInterpretation),
        "modality": str(getattr(first, "Modality", "")),
    })


def load_sequence(record):
    """Return full, finite float32 [D,H,W] voxels plus geometry, never labels/masks."""
    if record["kind"] == "dicom":
        return _load_dicom(record)
    if record["kind"] != "nifti":
        raise ValueError(f"Unsupported image kind: {record['kind']}")
    import nibabel as nib
    path = Path(record["path"])
    if path.name != f"{record['series_uid']}.nii.gz":
        raise ValueError(f"Only exact <SeriesUid>.nii.gz images are accepted: {path}")
    image = nib.load(str(path))
    if image.header.get_intent()[0] in {"label", "vector", "displacement vector", "rgb vector", "rgba vector"}:
        raise ValueError(f"{path}: annotation/vector NIfTI cannot be an image input")
    units = image.header.get_xyzt_units()[0]
    if units not in {"unknown", "mm", "meter", "micron"}:
        raise ValueError(f"{path}: unsupported spatial units {units}")
    original_affine = image.affine.copy()
    # NIfTI spatial units are normalized to mm, the native unit of DICOM geometry.
    scale = {"unknown": 1., "mm": 1., "meter": 1000., "micron": .001}[units]
    if scale != 1:
        scaled_affine = image.affine.copy()
        scaled_affine[:3, :] *= scale
        image = nib.Nifti1Image(image.get_fdata(dtype="float32"), scaled_affine)
    return _canonical_array(image, original_affine, image.shape, {
        "kind": "nifti", "source_path": str(path.resolve()),
        "original_spatial_units": units, "spatial_units": "mm",
        "spatial_units_assumed": units == "unknown",
    })


def discover_cases(dataset_path, cfg):
    """Index all exact named NIfTI or header-identified DICOM series and validate voxels."""
    root = Path(dataset_path)
    if not root.exists():
        raise FileNotFoundError(root)
    files = sorted(root.rglob("*") if root.is_dir() else [root])
    nifti_records, dicom_records = [], {}
    for path in files:
        if not path.is_file():
            continue
        if path.name == f"{path.parent.name}.nii.gz":
            nifti_records.append({
                "accession": _text_id(path.parent.parent.name, str(path)),
                "series_uid": _text_id(path.parent.name, str(path)),
                "patient_id": None, "kind": "nifti", "path": str(path.resolve()), "label": None,
            })
            continue
        if path.name.lower().endswith((".nii", ".nii.gz")):
            continue  # Prepared masks never enter the image index.
        candidate = path.suffix.lower() in {".dcm", ".dicom"} or not path.suffix
        if not candidate:
            # DICOM sometimes uses vendor-specific suffixes; inspect its standard magic.
            with path.open("rb") as stream:
                stream.seek(128)
                candidate = stream.read(4) == b"DICM"
        if not candidate:
            continue
        if not path.suffix:
            # Extensionless sidecar files are not necessarily DICOM.
            with path.open("rb") as stream:
                stream.seek(128)
                if stream.read(4) != b"DICM":
                    continue
        header = _dicom_header(path)
        if header is None:
            continue
        accession = _text_id(header.AccessionNumber, str(path))
        uid = _text_id(header.SeriesInstanceUID, str(path))
        key = (accession, uid)
        record = dicom_records.setdefault(key, {"accession": accession, "series_uid": uid,
                                              "patient_id": None, "kind": "dicom", "files": [], "label": None})
        record["files"].append(str(path.resolve()))
    index = {}
    accession_locations = {}
    for record in nifti_records + list(dicom_records.values()):
        if record["kind"] == "nifti":
            location = str(Path(record["path"]).parent.parent)
            previous = accession_locations.setdefault(record["accession"], location)
            if previous != location:
                raise ValueError(f"Accession occurs in multiple source folders: {record['accession']}")
        key = (record["accession"], record["series_uid"])
        if key in index:
            raise ValueError(f"Conflicting image sources for accession/series {key}; resolve aliases explicitly")
        _, metadata = load_sequence(record)  # Force complete decompression and voxel validation.
        record["patient_id"] = metadata.get("patient_id")
        index[key] = record
    if not index:
        raise ValueError(f"{root}: no exact named NIfTI or supported DICOM image series")
    return [index[key] for key in sorted(index)]


def _pair_rows(path, default_label, expected_label=None):
    pairs = []
    for number, row in enumerate(_table(path), 2):
        context = f"{path}:{number}"
        a = _text_id(_cell(row, "a", context), context)
        b = _text_id(_cell(row, "b", context), context)
        label = _cell(row, "label", context, required=False)
        has_label_column = any(key in row for key in _COLUMNS["label"])
        if has_label_column and label in (None, ""):
            continue  # An explicitly blank label is unknown, not a positive example.
        label = default_label if not has_label_column else label
        if str(label).strip() not in {"0", "1"}:
            raise ValueError(f"{context}: duplicate labels must be 0 or 1")
        label = int(label)
        if expected_label is not None and label != expected_label:
            raise ValueError(f"{context}: negative-pair file must contain only label 0")
        if a == b:
            raise ValueError(f"{context}: self-pairs are not accepted")
        a, b = sorted((a, b))
        pairs.append({"a": a, "b": b, "label": label})
    return pairs


def _components(accessions, pairs, patient_ids):
    parent = {accession: accession for accession in accessions}

    def find(value):
        while parent[value] != value:
            parent[value] = parent[parent[value]]
            value = parent[value]
        return value

    def union(a, b):
        a, b = find(a), find(b)
        if a != b:
            parent[max(a, b)] = min(a, b)

    for pair in pairs:
        if pair["label"] == 1:
            union(pair["a"], pair["b"])
    first_by_patient = {}
    for accession, patient in sorted(patient_ids.items()):
        if patient:
            union(accession, first_by_patient.setdefault(patient, accession))
    groups = defaultdict(list)
    for accession in sorted(accessions):
        groups[find(accession)].append(accession)
    return list(groups.values())


def _fixed_split(path, accessions):
    path = Path(path)
    result = {}

    def add(accession, split):
        accession = _text_id(accession, str(path))
        if split not in SPLIT_NAMES:
            raise ValueError(f"{path}: unknown split {split!r}")
        if accession in result:
            raise ValueError(f"{path}: repeated accession {accession}")
        result[accession] = split

    if path.suffix.lower() == ".json":
        obj = json.loads(path.read_text(encoding="utf-8"))
        if not isinstance(obj, dict):
            raise ValueError(f"{path}: split JSON must be an object")
        if any(key in obj for key in SPLIT_NAMES):
            if any(key not in {*SPLIT_NAMES, "meta"} for key in obj):
                raise ValueError(f"{path}: unknown keys in fixed split")
            for split in SPLIT_NAMES:
                values = obj.get(split)
                if not isinstance(values, list):
                    raise ValueError(f"{path}: a complete train/val/calibration split is required")
                for accession in values:
                    add(accession, split)
        else:
            for accession, split in obj.items():
                add(accession, split)
    else:
        for row in _table(path):
            add(_cell(row, "accession", str(path)), _cell(row, "split", str(path)))
    if set(result) != set(accessions):
        raise ValueError(f"{path}: fixed split must cover exactly all accessions; "
                         f"missing={sorted(set(accessions) - set(result))[:10]}, "
                         f"unknown={sorted(set(result) - set(accessions))[:10]}")
    return result


def _make_splits(records, pairs, settings, patient_ids):
    accessions = sorted({record["accession"] for record in records})
    groups = _components(accessions, pairs, patient_ids)
    if settings.get("split_file"):
        result = _fixed_split(settings["split_file"], accessions)
    else:
        ratios = settings.get("ratios", [.8, .1, .1])
        if (len(ratios) != 3 or any(not math.isfinite(float(v)) or float(v) <= 0 for v in ratios)
                or not math.isclose(sum(ratios), 1., abs_tol=1e-6)):
            raise ValueError("detection.data.ratios must be three positive numbers summing to one")
        labels_by_accession = defaultdict(Counter)
        for record in records:
            if record["label"] is not None:
                labels_by_accession[record["accession"]][record["label"]] += 1
        positive_counts = Counter(pair["a"] for pair in pairs if pair["label"] == 1)
        features = []
        for group in groups:
            counts = [len(group)] + [sum(labels_by_accession[a][label] for a in group) for label in range(3)]
            features.append((group, counts + [sum(positive_counts[a] for a in group)]))
        random.Random(int(settings.get("seed", 42))).shuffle(features)
        features.sort(key=lambda item: -len(item[0]))
        totals = [sum(item[1][i] for item in features) for i in range(5)]
        targets = [[total * float(ratio) for total in totals] for ratio in ratios]
        allocated = [[0] * 5 for _ in SPLIT_NAMES]
        result = {}
        for group, counts in features:
            def cost(index):
                return sum(((allocated[index][i] + counts[i] - targets[index][i]) ** 2
                            - (allocated[index][i] - targets[index][i]) ** 2)
                           / max(targets[index][i], 1.) for i in range(5))
            chosen = min(range(3), key=cost)
            allocated[chosen] = [a + b for a, b in zip(allocated[chosen], counts)]
            result.update({accession: SPLIT_NAMES[chosen] for accession in group})
    for group in groups:
        if len({result[accession] for accession in group}) != 1:
            raise ValueError(f"Fixed split leaks a patient or positive-pair component: {group[:10]}")
    return result


def build_manifest(cfg):
    """Build a validated, content-fingerprinted manifest without inventing negatives."""
    settings = dict(cfg["detection"]["data"])
    if not settings.get("split_file"):
        roots = {Path(cfg["paths"]["annotation_root"]), Path(cfg["paths"]["labels_dir"])}
        choices = sorted({root / name for root in roots for name in
                          ("official_split.csv", "splits.csv", "split.csv", "split.json")
                          if (root / name).is_file()})
        if len(choices) > 1:
            raise ValueError(f"Multiple possible official splits; configure split_file explicitly: {choices}")
        if choices:
            settings["split_file"] = str(choices[0])
    policy = settings.get("negative_policy", "unknown")
    if policy not in {"unknown", "explicit", "closed_world"}:
        raise ValueError("negative_policy must be unknown, explicit, or closed_world")
    warnings = []
    records = discover_cases(cfg["paths"]["annotation_root"], cfg)
    index = {(record["accession"], record["series_uid"]): record for record in records}
    accessions = {record["accession"] for record in records}
    labels_dir = Path(cfg["paths"]["labels_dir"])
    label_path = _label_file(labels_dir, "1_abnormal")
    source_tables = [label_path]
    label_patient_ids = {}
    for number, row in enumerate(_table(label_path), 2):
        context = f"{label_path}:{number}"
        accession = _text_id(_cell(row, "accession", context), context)
        uid = _text_id(_cell(row, "series_uid", context), context)
        value = str(_cell(row, "label", context)).strip().lower()
        if value not in LABEL_MAP:
            raise ValueError(f"{context}: unknown three-class label {value!r}")
        if (accession, uid) not in index:
            raise FileNotFoundError(f"{context}: no indexed image for {(accession, uid)}")
        record = index[(accession, uid)]
        if record["label"] is not None and record["label"] != LABEL_MAP[value]:
            raise ValueError(f"{context}: conflicting series labels")
        record["label"] = LABEL_MAP[value]
        patient = _cell(row, "patient_id", context, required=False)
        if patient not in (None, ""):
            patient = _text_id(patient, context)
            if accession in label_patient_ids and label_patient_ids[accession] != patient:
                raise ValueError(f"{context}: conflicting patient IDs")
            label_patient_ids[accession] = patient
    missing_labels = sum(record["label"] is None for record in records)
    if missing_labels:
        warnings.append(f"{missing_labels} image series have no three-class label; excluded from classification training")
    patient_ids = {accession: None for accession in accessions}

    def set_patient(accession, patient):
        if accession not in accessions:
            raise ValueError(f"Patient metadata references unknown accession {accession}")
        if patient_ids[accession] and patient_ids[accession] != patient:
            raise ValueError(f"Conflicting PatientID for accession {accession}")
        patient_ids[accession] = patient

    for record in records:
        if record["patient_id"]:
            set_patient(record["accession"], record["patient_id"])
    for accession, patient in label_patient_ids.items():
        set_patient(accession, patient)
    if settings.get("patient_metadata_file"):
        source_tables.append(Path(settings["patient_metadata_file"]))
        for row in _table(settings["patient_metadata_file"]):
            accession = _text_id(_cell(row, "accession", "patient metadata"), "patient metadata")
            patient = _text_id(_cell(row, "patient_id", "patient metadata"), "patient metadata")
            set_patient(accession, patient)
    for record in records:
        record["patient_id"] = patient_ids[record["accession"]]
    missing_patient = sum(patient is None for patient in patient_ids.values())
    if missing_patient:
        warnings.append(f"PatientID unavailable for {missing_patient} accessions; patient-level leakage cannot be fully verified")
    pair_path = _label_file(labels_dir, "2_duplicate")
    source_tables.append(pair_path)
    pairs = _pair_rows(pair_path, 1)
    if settings.get("negative_pair_file"):
        source_tables.append(Path(settings["negative_pair_file"]))
        pairs += _pair_rows(settings["negative_pair_file"], 0, expected_label=0)
    elif policy == "explicit" and not any(pair["label"] == 0 for pair in pairs):
        warnings.append("Explicit negative policy has no annotated negative pairs")
    pair_index = {}
    for pair in pairs:
        if pair["a"] not in accessions or pair["b"] not in accessions:
            raise FileNotFoundError(f"Pair references accession without images: {pair}")
        key = (pair["a"], pair["b"])
        if key in pair_index and pair_index[key]["label"] != pair["label"]:
            raise ValueError(f"Conflicting labels for pair {key}")
        pair_index[key] = pair
    pairs = [pair_index[key] for key in sorted(pair_index)]
    splits = _make_splits(records, pairs, settings, patient_ids)
    if policy == "closed_world":
        universe_path = settings.get("closed_world_ids_file")
        if not universe_path:
            raise ValueError("closed_world requires an explicit closed_world_ids_file")
        source_tables.append(Path(universe_path))
        universe = {_text_id(_cell(row, "accession", "closed-world universe"), "closed-world universe")
                    for row in _table(universe_path)}
        if not universe or not universe <= accessions:
            raise ValueError("Closed-world universe must be nonempty and contain only indexed accessions")
        # ponytail: exhaustive pairs are quadratic; use an explicit negative-pair file for large cohorts.
        for a, b in itertools.combinations(sorted(universe), 2):
            if (a, b) not in pair_index and splits[a] == splits[b]:
                pair_index[(a, b)] = {"a": a, "b": b, "label": 0}
        pairs = [pair_index[key] for key in sorted(pair_index)]
        warnings.append("Closed-world negatives are assumed only within the explicitly declared accession universe")
    if settings.get("split_file"):
        source_tables.append(Path(settings["split_file"]))
    duplicate_ready = policy != "unknown"
    if policy == "unknown":
        warnings.append("Duplicate supervision is NOT READY: unlisted pairs are unknown, never implicit negatives")
    coverage = {}
    for split in SPLIT_NAMES:
        class_counts = Counter(record["label"] for record in records
                               if splits[record["accession"]] == split and record["label"] is not None)
        pair_counts = Counter(pair["label"] for pair in pairs
                              if splits[pair["a"]] == split and splits[pair["b"]] == split)
        coverage[split] = {"accessions": sum(value == split for value in splits.values()),
                           "series_by_class": {str(label): class_counts[label] for label in range(3)},
                           "pairs_by_label": {str(label): pair_counts[label] for label in range(2)}}
        if set(class_counts) != {0, 1, 2}:
            warnings.append(f"{split}: incomplete three-class coverage {dict(class_counts)}")
        if set(pair_counts) != {0, 1}:
            duplicate_ready = False
            warnings.append(f"{split}: duplicate positive/negative coverage incomplete {dict(pair_counts)}")
    cross_split = sum(splits[pair["a"]] != splits[pair["b"]] for pair in pairs)
    if cross_split:
        warnings.append(f"{cross_split} annotated negative pairs cross splits and must be excluded from each split's training/evaluation")
    for record in records:
        paths = [record["path"]] if record["kind"] == "nifti" else record["files"]
        record["content_sha256"] = {path: file_hash(path) for path in paths}
    manifest = {"format_version": 1, "data_settings": cfg["detection"]["data"],
                "resolved_split_file": settings.get("split_file"), "records": records, "pairs": pairs,
                "splits": splits, "duplicate_ready": duplicate_ready,
                "negative_policy": policy, "coverage": coverage, "warnings": warnings,
                "source_tables": {str(path.resolve()): file_hash(path) for path in source_tables}}
    manifest["fingerprint"] = fingerprint(manifest)
    return manifest
